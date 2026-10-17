"""Write / merge the running 'Lead List' Excel file.

Existing rows are never overwritten (your Y/N answers, notes and any drop-downs you changed are kept).
New addresses from a scan are appended, then the whole list is re-sorted by score.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from .google import map_link
from .scoring import DEFAULT_SCORING, FACTOR_ORDER, score as calc_score

MAIN = ["Address (click = fence view)", "Map", "SCORE", "Action", "Service", "What to check",
        "Address confirmed?", "Fence matches?", "Still needs work?", "Talked to owner?",
        "Door hanger left?", "Outcome", "Notes"]
USER_COLS = MAIN[6:]  # filled in by you on the drive
TAIL = FACTOR_ORDER + ["Side", "Area / Alley", "Date scouted"]
OUTCOMES = "Quote booked,Interested,Not interested,No answer,Already done,Wrong address"
HEADER_ROW = 4

F = "Arial"
_thin = Side(style="thin", color="BBBBBB")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
HFILL = PatternFill("solid", fgColor="2F4F4F")
DFILL = PatternFill("solid", fgColor="808080")
YEL = PatternFill("solid", fgColor="FFF6CC")
BLUE = PatternFill("solid", fgColor="DDEBF7")


def _h(ws, r, c, v, fill=HFILL):
    x = ws.cell(r, c, v)
    x.font = Font(name=F, bold=True, color="FFFFFF")
    x.fill = fill
    x.alignment = Alignment(wrap_text=True, vertical="center")
    x.border = BORDER


def read_existing(path: Path) -> list[dict]:
    if not path.exists():
        return []
    wb = load_workbook(path)
    if "Lead List" not in wb.sheetnames:
        return []
    ws = wb["Lead List"]
    headers = [c.value for c in ws[HEADER_ROW]]
    rows = []
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        a = ws.cell(r, 1)
        if not a.value:
            continue
        rec = {"address": a.value, "link": a.hyperlink.target if a.hyperlink else ""}
        for i, h in enumerate(headers, 1):
            if h in (None, MAIN[0], "Map", "SCORE", "Action"):
                continue
            rec[h] = ws.cell(r, i).value
        rows.append(rec)
    return rows


def merge(existing: list[dict], new: list[dict]) -> tuple[list[dict], int]:
    seen = {r["address"].strip().lower() for r in existing}
    added = 0
    out = list(existing)
    for n in new:
        k = n["address"].strip().lower()
        if k in seen:
            continue
        seen.add(k)
        out.append(n)
        added += 1
    return out, added


def write_workbook(rows: list[dict], path: Path, cfg: dict = DEFAULT_SCORING) -> None:
    rows = sorted(rows, key=lambda r: -calc_score({f: r.get(f) for f in FACTOR_ORDER}, cfg))
    wb = Workbook()
    ws = wb.active
    ws.title = "Lead List"
    key = wb.create_sheet("Score Key")

    # ---- Score Key ----
    key["A1"] = "Fence Score - how points add up (max 100)"
    key["A1"].font = Font(name=F, bold=True, size=14)
    key["A2"] = ("Blue values are editable; every score on Lead List updates. "
                 "Starting guesses - adjust once you see which fences become jobs.")
    key["A2"].font = Font(name=F, size=9, italic=True)
    for c, h in enumerate(["Factor", "Level", "Points", "What it means"], 1):
        _h(key, 4, c, h)
    r = 4
    ranges = {}
    for name in FACTOR_ORDER:
        start = r + 1
        for i, (lvl, pts) in enumerate(cfg["factors"][name].items()):
            r += 1
            key.cell(r, 1, name if i == 0 else "")
            key.cell(r, 2, lvl)
            key.cell(r, 3, pts).fill = BLUE
            key.cell(r, 4, cfg["descriptions"].get(name, "") if i == 0 else "")
            for c in range(1, 5):
                key.cell(r, c).border = BORDER
                key.cell(r, c).font = Font(name=F, size=10, bold=(c == 1), color="0000FF" if c == 3 else "000000")
        ranges[name] = (start, r)
    r += 2
    key.cell(r, 1, "Go first at or above")
    key.cell(r, 3, cfg["go_first_at"]).fill = BLUE
    key.cell(r + 1, 1, "Worth a knock at or above")
    key.cell(r + 1, 3, cfg["worth_a_knock_at"]).fill = BLUE
    key.cell(r + 2, 1, "Below that = Skip (grayed out)")
    for rr in (r, r + 1, r + 2):
        key.cell(rr, 1).font = Font(name=F, size=10, italic=(rr == r + 2))
        key.cell(rr, 3).font = Font(name=F, size=10, color="0000FF")
    go, kn = f"'Score Key'!$C${r}", f"'Score Key'!$C${r + 1}"
    for c, w in zip("ABCD", [26, 12, 8, 44]):
        key.column_dimensions[c].width = w

    # ---- Lead List ----
    ws["A1"] = "Fence Scout - Lead List"
    ws["A1"].font = Font(name=F, bold=True, size=14)
    ws["A2"] = ("Click address -> Street View of the fence. Map -> house pin. Yellow = fill in. "
                "Scoring detail is at the far right (blue drop-downs).")
    ws["A2"].font = Font(name=F, size=9, color="666666")
    for i, h in enumerate(MAIN, 1):
        _h(ws, HEADER_ROW, i, h)
    for i, h in enumerate(TAIL, len(MAIN) + 1):
        _h(ws, HEADER_ROW, i, h, DFILL)
    ws.cell(HEADER_ROW - 1, len(MAIN) + 1, "Scoring detail ->").font = Font(name=F, size=9, bold=True, color="808080")

    dvs = []
    for name in FACTOR_ORDER:
        a, b = ranges[name]
        dv = DataValidation(type="list", formula1=f"'Score Key'!$B${a}:$B${b}")
        ws.add_data_validation(dv)
        dvs.append(dv)
    yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    oc = DataValidation(type="list", formula1=f'"{OUTCOMES}"', allow_blank=True)
    ws.add_data_validation(yn)
    ws.add_data_validation(oc)

    fc = len(MAIN) + 1
    ncols = len(MAIN) + len(TAIL)
    for j, rec in enumerate(rows):
        r = HEADER_ROW + 1 + j
        for c in range(1, ncols + 1):
            x = ws.cell(r, c)
            x.border = BORDER
            x.font = Font(name=F, size=10)
            x.alignment = Alignment(wrap_text=True, vertical="top")
        a = ws.cell(r, 1, rec["address"])
        if rec.get("link"):
            a.hyperlink = rec["link"]
            a.font = Font(name=F, size=10, color="0563C1", underline="single")
        m = ws.cell(r, 2, "Map")
        m.hyperlink = map_link(rec["address"])
        m.font = Font(name=F, size=10, color="0563C1", underline="single")
        parts = []
        for k, name in enumerate(FACTOR_ORDER):
            col = L(fc + k)
            c = ws[f"{col}{r}"]
            c.value = rec.get(name)
            c.fill = BLUE
            c.font = Font(name=F, size=10, color="0000FF")
            dvs[k].add(c.coordinate)
            s, e = ranges[name]
            parts.append(f"IFERROR(INDEX('Score Key'!$C${s}:$C${e},MATCH({col}{r},'Score Key'!$B${s}:$B${e},0)),0)")
        sc = ws.cell(r, 3, "=MIN(100," + "+".join(parts) + ")")
        sc.font = Font(name=F, size=12, bold=True)
        sc.alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(r, 4, f'=IF(C{r}>={go},"Go first",IF(C{r}>={kn},"Worth a knock","Skip"))').font = Font(name=F, size=10, bold=True)
        ws.cell(r, 5, rec.get("Service"))
        ws.cell(r, 6, rec.get("What to check"))
        for i, h in enumerate(USER_COLS, 7):
            ws.cell(r, i, rec.get(h)).fill = YEL
        for c in range(7, 12):
            yn.add(f"{L(c)}{r}")
        oc.add(f"L{r}")
        for i, h in enumerate(["Side", "Area / Alley", "Date scouted"]):
            ws.cell(r, fc + len(FACTOR_ORDER) + i, rec.get(h))

    last = max(HEADER_ROW + len(rows), HEADER_ROW + 1)
    endc = L(ncols)
    first = HEADER_ROW + 1
    ws.conditional_formatting.add(f"A{first}:{endc}{last}", FormulaRule(formula=[f'$D{first}="Skip"'], font=Font(color="A0A0A0")))
    ws.conditional_formatting.add(f"C{first}:D{last}", FormulaRule(formula=[f"$C{first}>={go}"], fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(f"C{first}:D{last}", FormulaRule(formula=[f"AND($C{first}>={kn},$C{first}<{go})"], fill=PatternFill("solid", fgColor="FFEB9C")))
    widths = [34, 6, 8, 13, 20, 36, 10, 9, 9, 9, 9, 15, 24, 10, 10, 9, 9, 9, 9, 10, 6, 30, 11]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[L(i)].width = w
    ws.freeze_panes = f"B{first}"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{endc}{last}"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_area = f"A1:M{last}"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
