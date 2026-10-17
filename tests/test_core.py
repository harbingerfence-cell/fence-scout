from datetime import date

import pytest
from openpyxl import load_workbook

from fence_scout.cache import NullCache
from fence_scout.export import merge, read_existing, write_workbook
from fence_scout.geo import LatLng, bearing_deg, haversine_m, offset, sample_polyline
from fence_scout.google import Pano
from fence_scout.osm import Way, build_query, parse_overpass
from fence_scout.pipeline import compass, plan_scan, run_scan
from fence_scout.scoring import DEFAULT_SCORING, action, imagery_age_level, score, service
from fence_scout.vision import Verdict, parse_verdict

A = LatLng(33.06428, -96.85285)   # south end of the Clear Well alley
B = LatLng(33.06620, -96.85278)   # north end


def test_geo_basics():
    d = haversine_m(A, B)
    assert 205 < d < 220
    assert bearing_deg(A, B) < 5 or bearing_deg(A, B) > 355
    east = offset(A, 90, 25)
    assert abs(haversine_m(A, east) - 25) < 0.1
    assert 85 < bearing_deg(A, east) < 95


def test_sampling_spacing():
    pts = sample_polyline([A, B], 12)
    assert len(pts) == int(haversine_m(A, B) // 12) + 1
    for p, q in zip(pts, pts[1:]):
        assert abs(haversine_m(p.point, q.point) - 12) < 0.5


def test_scoring_matches_spreadsheet_test():
    # 6549 Crown Forest from the manual test scored 58
    lv = {"Graying": "Heavy", "Stain fade": "Heavy", "New, unsealed": "No", "Damage": "None",
          "Run length": "Long", "Address conf.": "Medium", "Imagery age": "12-24 mo"}
    assert score(lv) == 58
    assert action(58) == "Go first"
    assert action(31) == "Worth a knock"
    assert action(18) == "Skip"
    assert service(lv) == "Stain"
    assert service({**lv, "Damage": "Minor"}) == "Repair ($99), then stain"


def test_imagery_age():
    t = date(2026, 10, 16)
    assert imagery_age_level("2025-04", t) == "12-24 mo"
    assert imagery_age_level("2026-03", t) == "<12 mo"
    assert imagery_age_level("2022-01", t) == ">24 mo"
    assert imagery_age_level("", t) == ">24 mo"


def test_parse_verdict_tolerant():
    v = parse_verdict('Here you go: {"fence_visible":"yes","material":"Wood","graying":"heavy",'
                      '"stain_fade":"Moderate","new_unsealed":"no","damage":"none","run_length":"long",'
                      '"house_number":"6605","notes":"gray boards","confidence":0.8}')
    assert v.is_wood_fence and v.graying == "Heavy" and v.material == "wood"
    assert v.house_number == "6605" and v.confidence == 0.8
    assert not parse_verdict("no json").is_wood_fence


def test_overpass_parse_and_query():
    data = {"elements": [
        {"type": "node", "id": 1, "lat": A.lat, "lon": A.lng},
        {"type": "node", "id": 2, "lat": B.lat, "lon": B.lng},
        {"type": "way", "id": 9, "nodes": [1, 2], "tags": {"highway": "service", "service": "alley"}},
    ]}
    ways = parse_overpass(data)
    assert ways[0].kind == "alley" and len(ways[0].coords) == 2
    assert '"service"="alley"' in build_query(A, 300, "alleys")
    assert "residential" in build_query(A, 300, "both")


def test_compass():
    assert compass(270) == "West" and compass(88) == "East" and compass(5) == "North"


# ---------- end-to-end with fakes ----------
class FakeGoogle:
    def __init__(self):
        self.paid_images = 0
        self.paid_geocodes = 0

    def pano_near(self, p, radius_m=10):
        # one pano every ~24 m along the alley (pano ids repeat for nearby samples)
        idx = round((p.lat - A.lat) / 0.000216)
        return Pano(f"pano{idx}", LatLng(A.lat + idx * 0.000216, A.lng), "2025-04")

    def image_b64(self, pano_id, heading):
        self.paid_images += 1
        return "x"

    def reverse_geocode(self, p):
        self.paid_geocodes += 1
        idx = round((p.lat - A.lat) / 0.000216)
        if p.lng < A.lng:
            return {"address": f"{6540 + idx * 4} Village Springs Dr, Plano, TX 75024", "location_type": "ROOFTOP",
                    "lat": p.lat, "lng": p.lng - 0.0001}
        return {"address": f"{6541 + idx * 4} Crown Forest Dr, Plano, TX 75024", "location_type": "ROOFTOP",
                "lat": p.lat, "lng": p.lng + 0.0001}


class FakeGrader:
    def __init__(self):
        self.calls = 0
        self.cache = NullCache()

    def cached(self, key):
        return None

    def grade(self, key, img):
        self.calls += 1
        heavy = key.startswith("pano3")
        east = 45 < float(key.split("|")[1]) < 135
        if heavy and east:
            return Verdict("yes", "wood", "Heavy", "Heavy", "No", "None", "Long", None, "long gray run", 0.9)
        if key.startswith("pano5"):
            return Verdict("no")  # no fence visible
        return Verdict("yes", "wood", "None", "None", "No", "None", "Medium", None, "stained", 0.8)


def test_end_to_end(tmp_path):
    g, gr = FakeGoogle(), FakeGrader()
    way = Way(1, "alley", "", [A, B])
    plan = plan_scan(A, [way], 400, 12, g, log=lambda *_: None)
    assert plan.panos >= 8
    assert len(plan.views) == plan.panos * 2
    leads = run_scan(plan, g, gr, DEFAULT_SCORING, "Test alley", log=lambda *_: None, today=date(2026, 10, 16))
    assert leads[0]["address"].startswith("6553 Crown Forest")  # pano3 east = heavy gray
    assert leads[0]["_action"] == "Go first"
    assert leads[0]["Side"] == "East"
    assert "[Photo 2025-04]" in leads[0]["What to check"]
    assert len({r["address"] for r in leads}) == len(leads)  # deduped

    out = tmp_path / "leads.xlsx"
    write_workbook(leads, out)
    # user fills in a field, then a second scan merges without clobbering it
    wb = load_workbook(out)
    ws = wb["Lead List"]
    ws["G5"] = "Y"
    ws["M5"] = "talked to owner, wants quote"
    wb.save(out)
    existing = read_existing(out)
    rows, added = merge(existing, leads + [{**leads[0], "address": "9999 New St, Plano, TX 75024"}])
    assert added == 1
    write_workbook(rows, out)
    ws = load_workbook(out)["Lead List"]
    assert ws["A5"].value.startswith("6553 Crown Forest")
    assert ws["G5"].value == "Y" and ws["M5"].value == "talked to owner, wants quote"
    assert ws["A5"].hyperlink.target.startswith("https://www.google.com/maps/@?api=1&map_action=pano")
    assert ws["C4"].value == "SCORE"
